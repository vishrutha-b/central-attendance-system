import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_header_styled(doc, text, level):
    p = doc.add_heading(text, level=level)
    run = p.runs[0]
    if level == 1:
        run.font.size = Pt(17)
        run.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
        run.font.bold = True
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(30, 41, 59) # Slate 800
        run.font.bold = True
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
    elif level == 3:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(71, 85, 105) # Slate 600
        run.font.bold = True
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
    return p

def add_callout(doc, text, title=None, alert_type="note"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    bg_color = "F8FAFC"
    border_color = "3B82F6"
    if alert_type == "warning":
        bg_color = "FFFBEB"
        border_color = "F59E0B"
    elif alert_type == "success":
        bg_color = "F0FDF4"
        border_color = "10B981"
        
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=140)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    if title:
        r_title = p.add_run(f"{title}: ")
        r_title.bold = True
        r_title.font.size = Pt(10)
        r_title.font.color.rgb = RGBColor(30, 41, 59)
        
    r_text = p.add_run(text)
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(51, 65, 85)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def build_word_document(output_path):
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(10)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("Central Attendance Management System (CAMS)")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)
    
    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_before = Pt(0)
    subtitle_p.paragraph_format.space_after = Pt(12)
    run_sub = subtitle_p.add_run("Complete Technical Specification, File Architecture & API Directory Reference")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = RGBColor(100, 116, 139)
    
    add_callout(
        doc,
        "Version: 1.0.0  |  Tech Stack: Python 3.12+, FastAPI, PyMongo Async, Pydantic v2, APScheduler, SMTP  |  Database: MongoDB 7.0 (cams_dev)  |  Interactive Swagger UI: http://localhost:8000/docs",
        title="SYSTEM METADATA",
        alert_type="note"
    )
    
    # 1. Executive Summary
    add_header_styled(doc, "1. Executive Summary & Core Objectives", level=1)
    doc.add_paragraph(
        "Central Attendance Management System (CAMS) is an enterprise backend service that tracks in-office "
        "presence, manages employee leaves, calculates policy compliance, and delivers real-time notifications."
    )
    
    bullets = [
        ("Precision Time Tracking: ", "Supports multiple clock-in and clock-out sessions per day with automatic duration calculations."),
        ("Instant Exit Violation Alerts: ", "When an employee exits before fulfilling the mandatory 5.0 hours in a day, the system sends an official alert email to that person the exact minute they exit."),
        ("Weekly Policy Auditing: ", "Evaluates compliance across the work week (≥ 3 days/week, ≥ 15 hours/week, ≥ 5 hours/day average) and generates weekly violation notices."),
        ("Smart Leave Protection: ", "Approved employee leaves and company holidays automatically scale down required weekly thresholds, protecting employees on scheduled time-off.")
    ]
    for b_title, b_desc in bullets:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_before = Pt(1)
        bp.paragraph_format.space_after = Pt(2)
        r1 = bp.add_run(b_title)
        r1.bold = True
        r1.font.color.rgb = RGBColor(30, 41, 59)
        r2 = bp.add_run(b_desc)
        r2.font.color.rgb = RGBColor(71, 85, 105)
        
    # 2. Company Attendance Policy
    add_header_styled(doc, "2. Company Attendance Policy Specification", level=1)
    policy_rows = [
        ["Policy Rule", "Threshold", "Trigger / Action"],
        ["Mandatory Daily Presence", "≥ 5.0 hours / day", "Instant Alert: Sent immediately upon clock-out if deficient."],
        ["Mandatory Weekly Days", "≥ 3 days / week", "Weekly monitoring cycle flags deficient employees."],
        ["Mandatory Weekly Hours", "≥ 15.0 hours / week", "Weekly monitoring cycle flags deficient employees."],
        ["Daily Attended Average", "≥ 5.0 hours / day", "Weekly monitoring cycle flags deficient employees."],
        ["Approved Leave Days", "Full day excused", "Automatically reduces required weekly days & hours proportionally."],
        ["Company Holidays", "Full day excused", "Automatically reduces required weekly days & hours proportionally."]
    ]
    tbl_policy = doc.add_table(rows=len(policy_rows), cols=3)
    tbl_policy.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Inches(2.0), Inches(1.5), Inches(3.0)]
    for r_idx, row in enumerate(policy_rows):
        for c_idx, val in enumerate(row):
            c = tbl_policy.cell(r_idx, c_idx)
            c.width = widths[c_idx]
            set_cell_margins(c, top=80, bottom=80, left=100, right=100)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(val)
            if r_idx == 0:
                set_cell_background(c, "1E293B")
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(c, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(30, 41, 59)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    # 3. Directory & File-by-File Guide
    add_header_styled(doc, "3. Directory & File-by-File Guide", level=1)
    files_guide = [
        (".env", "Configuration", "Stores sensitive settings: MongoDB URI, JWT secrets, Gmail SMTP credentials, and policy thresholds."),
        ("pyproject.toml / requirements.txt", "Build & Dependencies", "Declares all project dependencies, packaging configuration, and Ruff/Pytest settings."),
        ("app/main.py", "App Entrypoint", "FastAPI app instance, lifespan context manager (DB connect/disconnect & index creation), and router mounting."),
        ("app/config.py", "Settings Management", "Pydantic BaseSettings class that loads and validates environment variables."),
        ("app/exceptions.py", "Custom Exceptions", "Standardized exception handlers (NotFound, Validation, Conflict, Unauthorized)."),
        ("app/logging_config.py", "Structured Logging", "Configures structured JSON & console logging for observability and security."),
        ("app/database/connection.py", "Database Lifecycle", "Manages AsyncMongoClient connection pool singleton."),
        ("app/database/indexes.py", "Database Indexes", "Defines and builds indexes across all 7 MongoDB collections."),
        ("app/api/dependencies.py", "Dependency Injection", "Injects repos, services, and handles JWT bearer token validation and RBAC."),
        ("app/api/v1/attendance.py", "Attendance Endpoints", "Handles /entry, /exit, and /quick-record with automated daily email triggers."),
        ("app/api/v1/leave.py", "Leave & Holidays", "Public employee leave request endpoint, admin approvals/rejections, and holiday registry."),
        ("app/api/v1/employees.py", "Employee Directory", "Admin-protected CRUD operations to create, search, update, and deactivate employees."),
        ("app/api/v1/auth.py", "Authentication", "Admin login (/login) generating JWT tokens and identity profile (/me)."),
        ("app/api/v1/reports.py", "Reporting", "Generates daily, weekly, and monthly attendance reports and downloadable CSV exports."),
        ("app/api/v1/admin.py", "Administration", "System audit logs, notification audit trail, and manual weekly audit trigger."),
        ("app/services/attendance_service.py", "Attendance Logic", "Session recording, duration math, and instant trigger for daily violation emails upon exit."),
        ("app/services/policy_service.py", "Policy Rules Engine", "Weekly policy auditor (3 days, 15h, 5h avg) with dynamic leave/holiday scaling."),
        ("app/services/notification_service.py", "Email Engine", "Renders multi-part MIME emails with executive HTML templates and delivers via SMTP."),
        ("app/services/leave_service.py", "Leave Management", "Date validation, overlap prevention, approval workflows, and excused day calculations."),
        ("app/templates/daily_violation.html", "Daily Alert Email", "Professional, responsive HTML email sent when an employee exits before logging 5.0 hours."),
        ("app/templates/weekly_violation.html", "Weekly Warning Email", "Branded HTML email sent for weekly shortfalls in days or total hours.")
    ]
    tbl_files = doc.add_table(rows=len(files_guide) + 1, cols=3)
    tbl_files.alignment = WD_TABLE_ALIGNMENT.CENTER
    fwidths = [Inches(2.2), Inches(1.5), Inches(2.8)]
    for idx, h in enumerate(["File / Path", "Component Layer", "Functional Purpose"]):
        c = tbl_files.cell(0, idx)
        c.width = fwidths[idx]
        set_cell_background(c, "1E293B")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
    for r_idx, (fp, fl, fd) in enumerate(files_guide):
        row = r_idx + 1
        for c_idx, text in enumerate([fp, fl, fd]):
            c = tbl_files.cell(row, c_idx)
            c.width = fwidths[c_idx]
            set_cell_margins(c, top=60, bottom=60, left=100, right=100)
            set_cell_background(c, "F8FAFC" if row % 2 == 1 else "FFFFFF")
            p = c.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(15, 23, 42)
            else:
                r.font.color.rgb = RGBColor(51, 65, 85)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 4. Comprehensive API Directory: What Each API Is Used For
    add_header_styled(doc, "4. Comprehensive API Directory: What Each API Is Used For", level=1)
    doc.add_paragraph(
        "This section details every API endpoint in the system, explaining who calls it, its core purpose, "
        "the key parameters it expects, and the automated side effects it triggers."
    )
    
    api_list = [
        # Method, Endpoint, Target Role, Purpose, Key Inputs, Side Effects
        ("POST", "/api/v1/attendance/entry", "Employee / Biometric Device", "Clock-In: Marks start of workday", "employee_id, entry_time ('09:00'), date", "Verifies active status; creates open session with status PRESENT."),
        ("POST", "/api/v1/attendance/exit", "Employee / Biometric Device", "Clock-Out: Marks departure from office", "employee_id, exit_time ('12:00'), date", "Calculates hours. IF < 5.0h, INSTANTLY sends Daily Violation Email to employee."),
        ("POST", "/api/v1/attendance/quick-record", "Biometric Sync / HR Tester", "Combined Clock-In & Clock-Out in one call", "employee_id, date, entry_time, exit_time", "Computes session hours; triggers email alert if < 5.0h."),
        ("GET", "/api/v1/attendance/{employee_id}", "Employee / Manager", "Live Attendance Status for today", "employee_id (path)", "Returns current open clock-in session or today's completed hours."),
        ("GET", "/api/v1/attendance/{employee_id}/history", "Employee / HR Admin", "Attendance Timesheet & History", "employee_id, start_date, end_date", "Returns chronological list of all sessions, timestamps, and durations."),
        ("PATCH", "/api/v1/attendance/session/{id}", "HR Admin (Auth Required)", "Attendance Regularization / Correction", "attendance_id, entry_time, exit_time, notes", "Updates session times and logs action in immutable audit_logs."),
        
        ("POST", "/api/v1/leave", "Employee (Self-Service)", "Apply for Leave (Sick, Casual, Annual)", "employee_id, leave_type, start_date, end_date, reason", "Validates dates; checks for duplicates; creates PENDING leave without admin login."),
        ("GET", "/api/v1/leave", "HR Admin / Manager", "List Employee Leave Applications", "employee_id, status (pending/approved), dates", "Returns filtered list of employee leave requests."),
        ("GET", "/api/v1/leave/{id}", "Employee / HR Admin", "Inspect Single Leave Details", "leave_id (path)", "Returns full leave record and review status."),
        ("PATCH", "/api/v1/leave/{id}/approve", "HR Admin (Auth Required)", "Approve Leave Request", "leave_id, optional review note", "Marks leave APPROVED; automatically reduces required weekly days & hours during audits."),
        ("PATCH", "/api/v1/leave/{id}/reject", "HR Admin (Auth Required)", "Reject Leave Request", "leave_id, rejection reason", "Updates status to REJECTED with recorded reviewer remarks."),
        ("POST", "/api/v1/holidays", "HR Admin (Auth Required)", "Declare Company Holiday", "name, date, description", "Registers holiday in holidays collection; automatically excused during weekly audits."),
        ("GET", "/api/v1/holidays", "All Employees / Admins", "Company Holiday Calendar", "year (optional query)", "Returns official holiday dates."),

        ("POST", "/api/v1/employees", "HR Admin (Auth Required)", "Onboard New Employee", "employee_id, name, email, department, role", "Ensures unique ID/email; saves employee and creates audit log."),
        ("GET", "/api/v1/employees", "HR Admin (Auth Required)", "Search & List Employee Directory", "search, department, is_active, page", "Returns paginated list of active/inactive employees."),
        ("GET", "/api/v1/employees/{id}", "HR Admin / Manager", "View Employee Profile", "employee_id (path)", "Returns full employee profile and department info."),
        ("PATCH", "/api/v1/employees/{id}", "HR Admin (Auth Required)", "Update Employee Profile", "name, email, department, role", "Updates employee and records change in audit logs."),
        ("DELETE", "/api/v1/employees/{id}", "HR Admin (Auth Required)", "Offboard Employee", "employee_id (path)", "Soft-deletes employee (is_active = false) preserving history."),

        ("GET", "/api/v1/reports/attendance", "HR Admin / Executives", "Attendance Summary & Compliance Analytics", "start_date, end_date, department", "Computes total days, hours, averages, and lists all non-compliant employees."),
        ("GET", "/api/v1/reports/attendance/export.csv", "HR / Payroll Administrator", "Download Attendance as CSV Spreadsheet", "start_date, end_date, department", "Streams raw CSV file for Excel, Google Sheets, or payroll systems."),

        ("POST", "/api/v1/auth/login", "Administrators", "Admin Login & Token Generation", "username, password", "Verifies Argon2 hash; issues signed JWT Bearer token valid 24h."),
        ("GET", "/api/v1/auth/me", "Authenticated Admin", "Admin Identity & Role Verification", "Bearer JWT Token", "Returns current administrator profile and permissions."),
        ("POST", "/api/v1/admin/users", "Super Administrator", "Create Additional Admin User", "username, email, password, role", "Hashes password and stores new admin in admin_users."),
        ("GET", "/api/v1/admin/audit-logs", "Compliance / Auditors", "Inspect Immutable System Audit Trail", "action, resource_type, dates", "Chronological audit trail of all administrative events."),
        ("POST", "/api/v1/admin/monitoring/run", "HR Admin / APScheduler Cron", "Trigger Weekly Compliance Audit On-Demand", "target_week_start, include_current_week, force_resend", "Runs policy engine, factors in leaves/holidays, and sends weekly violation emails."),

        ("GET", "/health/live", "DevOps / Docker Probe", "Service Liveness Probe", "None", "Returns 200 OK if FastAPI process is alive."),
        ("GET", "/health/ready", "DevOps / Docker Probe", "Database Readiness Probe", "None", "Pings MongoDB; returns 200 OK if DB connected or 503 if down.")
    ]
    
    tbl_api = doc.add_table(rows=len(api_list) + 1, cols=4)
    tbl_api.alignment = WD_TABLE_ALIGNMENT.CENTER
    api_widths = [Inches(1.8), Inches(1.3), Inches(1.6), Inches(1.8)]
    
    api_headers = ["Endpoint & Method", "Target User", "What It Is Used For", "Automated Action / Side Effects"]
    for idx, h in enumerate(api_headers):
        c = tbl_api.cell(0, idx)
        c.width = api_widths[idx]
        set_cell_background(c, "1E293B")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    for row_idx, (meth, ep, usr, desc, inps, sides) in enumerate(api_list):
        row = row_idx + 1
        c0 = tbl_api.cell(row, 0)
        c1 = tbl_api.cell(row, 1)
        c2 = tbl_api.cell(row, 2)
        c3 = tbl_api.cell(row, 3)
        
        for ci, cell in enumerate([c0, c1, c2, c3]):
            cell.width = api_widths[ci]
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            set_cell_background(cell, "F8FAFC" if row % 2 == 1 else "FFFFFF")
            
        p0 = c0.paragraphs[0]
        r_meth = p0.add_run(f"[{meth}] ")
        r_meth.bold = True
        r_meth.font.size = Pt(8.5)
        r_meth.font.color.rgb = RGBColor(37, 99, 235) if meth == "POST" else (RGBColor(16, 185, 129) if meth == "GET" else RGBColor(245, 158, 11))
        r_ep = p0.add_run(ep)
        r_ep.font.name = "Consolas"
        r_ep.font.size = Pt(8)
        r_ep.font.color.rgb = RGBColor(15, 23, 42)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(usr)
        r1.font.size = Pt(8.5)
        r1.font.color.rgb = RGBColor(51, 65, 85)
        
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(f"{desc}\nInputs: {inps}")
        r2.font.size = Pt(8)
        r2.font.color.rgb = RGBColor(30, 41, 59)
        
        p3 = c3.paragraphs[0]
        r3 = p3.add_run(sides)
        r3.font.size = Pt(8)
        r3.font.color.rgb = RGBColor(71, 85, 105)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    # 5. Real-Time Exit Alerts & Email Formatting
    add_header_styled(doc, "5. Real-Time Exit Alerts & Email Formatting", level=1)
    doc.add_paragraph(
        "When an employee clocks in and exits before fulfilling the mandatory 5.0 hours in a day, "
        "the attendance service immediately detects the shortfall during the exit transaction."
    )
    add_callout(
        doc,
        "The alert triggers instantaneously inside the record_exit() and quick_record() methods. "
        "No cron delay is required. The API response also includes a 'policy_alert' summary confirming dispatch.",
        title="INSTANT ENFORCEMENT",
        alert_type="warning"
    )
    
    # 6. Quality Assurance & Tests
    add_header_styled(doc, "6. Quality Assurance & Tests", level=1)
    doc.add_paragraph(
        "The project includes 39 automated tests covering authentication, employee lifecycle, "
        "instant exit email alerts, leave workflows, and compliance audits."
    )
    qa_points = [
        ("Automated Test Suite: ", "39 test cases (100% passing in ~1.2s)."),
        ("Static Analysis & Linting: ", "Ruff linting clean with 0 errors across app and tests."),
        ("Modern PyMongo Async: ", "Native AsyncMongoClient without deprecated frameworks.")
    ]
    for q_title, q_desc in qa_points:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(q_title)
        r1.bold = True
        r2 = p.add_run(q_desc)
        r2.font.color.rgb = RGBColor(51, 65, 85)
        
    doc.save(output_path)
    print(f"Successfully generated Word document: {output_path}")

if __name__ == "__main__":
    out_file = os.path.abspath("Central_Attendance_System_Documentation.docx")
    build_word_document(out_file)
