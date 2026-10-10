import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_header_styled(doc, text, level):
    p = doc.add_heading(text, level=level)
    run = p.runs[0]
    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(15, 23, 42)
        run.font.bold = True
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(30, 41, 59)
        run.font.bold = True
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
    elif level == 3:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(71, 85, 105)
        run.font.bold = True
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
    return p

def add_callout(doc, text, title=None, alert_type="note"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    bg_color = "F8FAFC"
    border_color = "3B82F6"
    if alert_type == "warning":
        bg_color = "FFFBEB"
        border_color = "F59E0B"
    elif alert_type == "success":
        bg_color = "F0FDF4"
        border_color = "10B981"
        
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    
    if title:
        run_title = p.add_run(f"[{title}] ")
        run_title.bold = True
        run_title.font.size = Pt(9.5)
        run_title.font.color.rgb = RGBColor(15, 23, 42)
        
    run_text = p.add_run(text)
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = RGBColor(51, 65, 85)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_docx(output_path):
    doc = Document()
    
    # Page setup
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)
        
    # Title Header Block
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("Central Attendance Management System (CAMS)")
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 58, 138)
    
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("CI/CD Pipeline Specification & DevSecOps Architecture Guide")
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(100, 116, 139)
    
    add_callout(
        doc,
        "This official technical documentation specifies every automated stage, tool, command, and security scan "
        "implemented in the GitHub Actions CI/CD pipeline (.github/workflows/ci.yml).",
        title="DOCUMENT PURPOSE",
        alert_type="note"
    )
    
    # 1. Executive Architecture
    add_header_styled(doc, "1. Executive Pipeline Architecture", level=1)
    p_arch = doc.add_paragraph()
    p_arch.add_run(
        "The automated CI/CD pipeline is designed as a Directed Acyclic Graph (DAG) multi-job tree. "
        "Instead of running as a single linear sequence, jobs branch out in parallel to optimize build "
        "velocity and aggregate into quality and security gates before final Docker containerization."
    )
    
    # DAG Tree Table
    tbl_tree = doc.add_table(rows=7, cols=3)
    tbl_tree.alignment = WD_TABLE_ALIGNMENT.CENTER
    tree_headers = ["Stage Node", "Job Identifier", "Dependencies & Flow"]
    tree_data = [
        ("Root Node", "syntax-and-setup", "Triggered on Push/PR -> Compiles Python AST"),
        ("Branch A", "code-quality", "needs: syntax-and-setup -> Runs Ruff & MyPy in parallel"),
        ("Branch B", "security-scans", "needs: syntax-and-setup -> Runs Bandit, pip-audit, Safety, Gitleaks"),
        ("Branch C", "test-suite", "needs: syntax-and-setup -> Spins MongoDB, executes Unit & API tests"),
        ("Merge Node", "sonarqube", "needs: [code-quality, security-scans, test-suite] -> Evaluates Quality Gate"),
        ("Final Node", "docker-build-and-scan", "needs: sonarqube -> Builds image & executes Aqua Trivy container scan")
    ]
    
    for idx, h in enumerate(tree_headers):
        c = tbl_tree.cell(0, idx)
        set_cell_background(c, "1E293B")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    for row_idx, (col0, col1, col2) in enumerate(tree_data):
        row = row_idx + 1
        for c_idx, val in enumerate([col0, col1, col2]):
            c = tbl_tree.cell(row, c_idx)
            set_cell_background(c, "F8FAFC" if row % 2 == 1 else "FFFFFF")
            set_cell_margins(c, top=60, bottom=60, left=100, right=100)
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            if c_idx == 1:
                r.font.name = "Consolas"
                r.bold = True
                
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # 2. Detailed Stages
    add_header_styled(doc, "2. Deep-Dive Stage Analysis", level=1)
    
    stages_info = [
        (
            "Stage 1: Syntax & Base Verification (syntax-and-setup)",
            "Root Gatekeeper",
            "Validates the entire codebase using Python's native compileall module. "
            "It scans app/, tests/, and scripts/ to ensure all files compile cleanly into bytecode.",
            "Fast-fails in seconds if missing commas, unmatched parentheses, or indentation errors exist before spinning up compute-heavy services.",
            "python -m compileall app/ tests/ scripts/"
        ),
        (
            "Stage 2: Code Quality & Type Checking (code-quality)",
            "Ruff & MyPy",
            "Applies linting standards, import organization, formatting checks, and static type analysis across the application code.",
            "Enforces PEP 8 compliance, cleans unused imports, and validates type consistency in FastAPI Pydantic models.",
            "ruff check app/ tests/ scripts/\nmypy --explicit-package-bases --ignore-missing-imports app/"
        ),
        (
            "Stage 3: Comprehensive Security Scans (security-scans)",
            "Bandit, pip-audit, Safety & Gitleaks",
            "Multi-vector security auditing: Static Application Security Testing (SAST) via Bandit, "
            "dependency CVE scans via pip-audit & Safety, and secret scanning via Gitleaks.",
            "Prevents OWASP Top 10 injection vulnerabilities, known library CVEs, and accidental commits of API keys or JWT tokens.",
            "bandit -r app/ -ll -s B104\npip-audit -r requirements.txt\nsafety check -r requirements.txt\ngitleaks detect"
        ),
        (
            "Stage 4: Automated Testing & Coverage (test-suite)",
            "Pytest, pytest-cov & MongoDB 7.0 Service Container",
            "Spins up an isolated MongoDB 7.0 Docker service with health checks, executes Unit tests (tests/unit) "
            "and API integration tests (tests/api), and publishes coverage reports.",
            "Ensures 100% functional test suite reliability, validates business policy rules, and exports coverage.xml for quality tracking.",
            "pytest tests/unit -v --cov=app --cov-report=xml:coverage.xml\npytest tests/api -v"
        ),
        (
            "Stage 5: SonarQube Quality Gate (sonarqube)",
            "SonarSource Scanner / SonarCloud",
            "Ingests coverage.xml and source files to evaluate duplication, code smells, test coverage, and security hotspots.",
            "Guarantees adherence to Clean Code standards. Automatically skips cleanly if SONAR_TOKEN secret is not yet configured.",
            "sonar-scanner (via sonar-project.properties)"
        ),
        (
            "Stage 6: Docker Container Build & Security Scan (docker-build-and-scan)",
            "Docker Buildx & Aqua Trivy",
            "Builds the production Docker image and runs an Aqua Trivy container layer vulnerability scan.",
            "Ensures final deployable container image contains zero unmitigated CRITICAL or HIGH operating system vulnerabilities.",
            "docker build -t central-attendance-system:$SHA .\ntrivy image --severity CRITICAL,HIGH central-attendance-system:$SHA"
        )
    ]
    
    for title, tool, desc, objective, cmd in stages_info:
        add_header_styled(doc, title, level=2)
        
        p = doc.add_paragraph()
        r1 = p.add_run("Tools Used: ")
        r1.bold = True
        r1.font.size = Pt(9.5)
        r2 = p.add_run(f"{tool}\n")
        r2.font.size = Pt(9.5)
        
        r3 = p.add_run("Objective: ")
        r3.bold = True
        r3.font.size = Pt(9.5)
        r4 = p.add_run(f"{objective}\n")
        r4.font.size = Pt(9.5)
        
        r5 = p.add_run("Description: ")
        r5.bold = True
        r5.font.size = Pt(9.5)
        r6 = p.add_run(f"{desc}")
        r6.font.size = Pt(9.5)
        
        p_cmd = doc.add_paragraph()
        p_cmd.paragraph_format.space_before = Pt(4)
        p_cmd.paragraph_format.space_after = Pt(8)
        p_cmd.paragraph_format.left_indent = Inches(0.2)
        r_cmd = p_cmd.add_run(f"Command:\n{cmd}")
        r_cmd.font.name = "Consolas"
        r_cmd.font.size = Pt(8.5)
        r_cmd.font.color.rgb = RGBColor(30, 41, 59)
        
    # 3. Quick Reference Matrix
    add_header_styled(doc, "3. Local Developer Verification Quick Guide", level=1)
    add_callout(
        doc,
        "Before pushing code to Git, developers can run this quick 4-step checklist locally in PowerShell:\n\n"
        "1. python -m compileall app/ tests/ scripts/\n"
        "2. ruff check app/ tests/ scripts/\n"
        "3. bandit -r app/ -ll -s B104\n"
        "4. pytest tests/unit tests/api -v",
        title="DEVELOPER CHEATSHEET",
        alert_type="success"
    )
    
    doc.save(output_path)
    print(f"Successfully generated Word document: {output_path}")

if __name__ == "__main__":
    out_file = Path("docs/CI_CD_PIPELINE_DOCUMENTATION.docx").resolve()
    build_docx(str(out_file))
